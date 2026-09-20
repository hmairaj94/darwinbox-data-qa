from __future__ import annotations

from numbers import Number
from typing import Any

import pandas as pd
from openpyxl.utils import get_column_letter


def _non_empty_mask(frame: pd.DataFrame) -> pd.DataFrame:
    mask = frame.notna()
    for column in frame.columns:
        strings = frame[column].map(lambda value: isinstance(value, str) and not value.strip())
        mask.loc[strings, column] = False
    return mask


def _segments(flags: list[bool]) -> list[tuple[int, int]]:
    segments: list[tuple[int, int]] = []
    start: int | None = None
    for index, populated in enumerate([*flags, False]):
        if populated and start is None:
            start = index
        elif not populated and start is not None:
            segments.append((start, index))
            start = None
    return segments


def _cell_metrics(values: list[Any]) -> tuple[float, float, float]:
    present = [value for value in values if pd.notna(value) and str(value).strip()]
    if not values or not present:
        return 0.0, 0.0, 0.0
    density = len(present) / len(values)
    text_ratio = sum(isinstance(value, str) for value in present) / len(present)
    unique_ratio = len({str(value).strip().lower() for value in present}) / len(present)
    return density, text_ratio, unique_ratio


def _header_score(region: pd.DataFrame, position: int) -> float:
    density, text_ratio, unique_ratio = _cell_metrics(region.iloc[position].tolist())
    following = region.iloc[position + 1 : position + 4]
    if following.empty:
        return 0.0
    following_mask = _non_empty_mask(following)
    following_density = float(following_mask.mean(axis=1).mean())
    following_values = [
        value
        for value in following.to_numpy().ravel().tolist()
        if pd.notna(value) and str(value).strip()
    ]
    numeric_ratio = (
        sum(isinstance(value, Number) and not isinstance(value, bool) for value in following_values)
        / len(following_values)
        if following_values
        else 0.0
    )
    type_transition = text_ratio * numeric_ratio
    return round(
        (0.3 * density)
        + (0.25 * text_ratio)
        + (0.15 * unique_ratio)
        + (0.2 * following_density)
        + (0.1 * type_transition),
        3,
    )


def _header_bounds(region: pd.DataFrame, best_position: int) -> tuple[int, int]:
    width = len(region.columns)
    start = best_position
    end = best_position
    if best_position > 0:
        previous = region.iloc[best_position - 1].tolist()
        density, text_ratio, _ = _cell_metrics(previous)
        populated = round(density * width)
        if 2 <= populated < width and text_ratio >= 0.8:
            start = best_position - 1
    current = region.iloc[best_position].tolist()
    density, text_ratio, _ = _cell_metrics(current)
    if best_position + 2 < len(region) and density < 1 and text_ratio >= 0.8:
        next_density, next_text_ratio, _ = _cell_metrics(
            region.iloc[best_position + 1].tolist()
        )
        if next_density >= 0.5 and next_text_ratio >= 0.7:
            end = best_position + 1
    return start, end


def _flatten_headers(header: pd.DataFrame) -> list[str]:
    filled = header.copy()
    for index in filled.index:
        filled.loc[index] = filled.loc[index].ffill()
    columns: list[str] = []
    for column_position in range(len(filled.columns)):
        parts: list[str] = []
        for value in filled.iloc[:, column_position].tolist():
            if pd.isna(value) or not str(value).strip():
                continue
            part = str(value).strip()
            if not parts or parts[-1].lower() != part.lower():
                parts.append(part)
        columns.append("_".join(parts) if parts else f"column_{column_position + 1}")
    return columns


def _source_range(
    row_start: int, row_end: int, column_start: int, column_end: int
) -> str:
    return (
        f"{get_column_letter(column_start + 1)}{row_start + 1}:"
        f"{get_column_letter(column_end)}{row_end}"
    )


def detect_sheet_tables(
    raw: pd.DataFrame,
    header_scan_rows: int = 12,
    min_table_rows: int = 2,
    confidence_threshold: float = 0.65,
) -> list[dict[str, Any]]:
    """Detect rectangular tables and their headers in a raw Excel sheet."""
    if raw.empty:
        return []
    mask = _non_empty_mask(raw)
    row_regions = _segments(mask.any(axis=1).tolist())
    detected: list[dict[str, Any]] = []

    for row_start, row_end in row_regions:
        row_block = raw.iloc[row_start:row_end]
        block_mask = _non_empty_mask(row_block)
        column_regions = _segments(block_mask.any(axis=0).tolist())
        for column_start, column_end in column_regions:
            region = raw.iloc[row_start:row_end, column_start:column_end].copy()
            if len(region) < min_table_rows or region.empty:
                continue
            candidates = range(min(header_scan_rows, len(region) - 1))
            scores = [(position, _header_score(region, position)) for position in candidates]
            if not scores:
                continue
            best_position, confidence = max(scores, key=lambda item: item[1])
            header_start, header_end = _header_bounds(region, best_position)
            data = region.iloc[header_end + 1 :].copy()
            data = data.dropna(axis=0, how="all").dropna(axis=1, how="all")
            if data.empty:
                continue
            header = region.iloc[header_start : header_end + 1, :]
            header_names = _flatten_headers(header)
            retained_positions = [region.columns.get_loc(column) for column in data.columns]
            data.columns = [header_names[position] for position in retained_positions]
            data = data.reset_index(drop=True)

            warnings: list[str] = []
            if header_start > 0:
                warnings.append(
                    f"Ignored {header_start} leading row(s) before the detected header."
                )
            if header_end > header_start:
                warnings.append("Flattened a multi-row header.")
            if confidence < confidence_threshold:
                warnings.append("Low-confidence header detection; verify the preview.")
            detected.append(
                {
                    "frame": data,
                    "source_range": _source_range(
                        row_start + header_start,
                        row_end,
                        column_start,
                        column_end,
                    ),
                    "detection_confidence": confidence,
                    "warnings": warnings,
                }
            )

    if len(detected) > 1:
        for table in detected:
            table["warnings"].append("Multiple table regions were detected in this sheet.")
    return detected

